#!/usr/bin/env python3
"""
Assemble les exports CSV du référentiel en un classeur unique et lisible.

TB « Exploration de l'IA pour les entreprises industrielles ».

Le classeur est une PROJECTION DATÉE de la base consolidée, jamais une
saisie manuelle : c'est ce qui garantit qu'il ne peut pas diverger du
référentiel ni de l'annexe 1 en Markdown, produits par la même source.

Exécution — depuis prototype/, sans rien installer sur le poste :

  bash exports/generer_csv.sh
  docker run --rm --user "$(id -u):$(id -g)" -e HOME=/tmp \
    -v "$PWD/exports:/w" -w /w python:3.12-slim \
    sh -c "pip install -q --target /tmp/libs openpyxl && PYTHONPATH=/tmp/libs python csv_vers_xlsx.py"
"""

import csv
import datetime
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

ICI = pathlib.Path(__file__).parent
CSV = ICI / "csv"
SORTIE = ICI / "1_tableau_de_veille.xlsx"

POLICE = "Arial"
ENCRE = "1C2733"
GRIS_FILET = "D3D6DA"

ENTETE = PatternFill("solid", fgColor=ENCRE)
POLICE_ENTETE = Font(name=POLICE, color="FFFFFF", bold=True, size=10)
POLICE_CORPS = Font(name=POLICE, size=10)
FILET = Side(style="thin", color=GRIS_FILET)
BORDURE = Border(left=FILET, right=FILET, top=FILET, bottom=FILET)

# Statuts colorés : le lecteur doit voir d'un coup d'œil ce qui est
# certifié et ce qui ne l'est pas. La nuance porte l'essentiel du propos.
VERT, AMBRE, ROUGE, GRIS = "D6EBD8", "FBF0D0", "F7D9D5", "F0F1F3"

COULEURS_STATUT = {
    "certifie": VERT,
    "certifiee": VERT,
    "a_confirmer": AMBRE,
    "restreint": ROUGE,
    "restreinte": ROUGE,
}

COULEURS_COUVERTURE = {
    "couverte": VERT,
    "couverte a confirmer": AMBRE,
    "non couverte": ROUGE,
}

COULEURS_CRITICITE = {
    "dominante": "E8E2F2",
    "significative": "F2F0F7",
    "marginale": GRIS,
}

# Colonnes larges par nature : sans plafond différencié, une colonne de
# mécanisme causal écrase toute la feuille.
LARGEURS_MAX = {
    "Mécanisme causal en jeu": 70,
    "Question telle qu'elle se pose": 60,
    "Question non couverte": 60,
    "Formulation générique": 60,
    "Formulation": 60,
    "Remarque": 45,
    "Indicateur": 50,
    "URL de la source": 40,
    "URL": 40,
}

# Ordre délibéré : le bilan d'abord (c'est le chiffre qui fait foi), les
# lacunes juste après les grilles (ce qui manque se lit aussi vite que ce
# qui existe), les sources en fin car ce sont des pièces justificatives.
FEUILLES = [
    ("bilan.csv", "Bilan"),
    ("indicateurs.csv", "Indicateurs"),
    ("instanciation_qv.csv", "Instanciation QV"),
    ("questions_de_veille.csv", "Questions de veille"),
    ("couverture_qv.csv", "Couverture QV"),
    ("lacunes.csv", "Lacunes"),
    ("sources.csv", "Sources"),
]


def nombre_si_possible(valeur):
    if valeur is None or valeur == "":
        return None
    try:
        return int(valeur)
    except ValueError:
        pass
    try:
        return float(valeur.replace(",", "."))
    except ValueError:
        return valeur


def index_colonne(entetes, *fragments):
    for i, entete in enumerate(entetes):
        if any(f.lower() in entete.lower() for f in fragments):
            return i + 1
    return None


def ajouter_feuille(classeur, chemin, titre):
    with open(chemin, encoding="utf-8", newline="") as f:
        lignes = list(csv.reader(f))
    if not lignes:
        return None
    entetes, donnees = lignes[0], lignes[1:]

    ws = classeur.create_sheet(titre)
    ws.append(entetes)
    for cellule in ws[1]:
        cellule.fill = ENTETE
        cellule.font = POLICE_ENTETE
        cellule.alignment = Alignment(
            vertical="center", horizontal="left", wrap_text=True
        )
    ws.row_dimensions[1].height = 32

    for ligne in donnees:
        ws.append([nombre_si_possible(v) for v in ligne])

    col_statut = index_colonne(entetes, "Statut")
    col_couverture = index_colonne(entetes, "Couverture", "Nature")
    col_criticite = index_colonne(entetes, "Criticité")
    col_obs = index_colonne(entetes, "Observations en base")

    for r in range(2, ws.max_row + 1):
        remplissage = None
        if col_statut:
            v = str(ws.cell(row=r, column=col_statut).value or "").strip()
            remplissage = COULEURS_STATUT.get(v)
        if remplissage is None and col_couverture:
            v = str(ws.cell(row=r, column=col_couverture).value or "").strip()
            remplissage = COULEURS_COUVERTURE.get(v)
            if remplissage is None and v.startswith("lacune"):
                remplissage = ROUGE
            elif remplissage is None and v.startswith("couverture nominale"):
                remplissage = AMBRE
        if remplissage is None and col_criticite:
            v = str(ws.cell(row=r, column=col_criticite).value or "").strip()
            remplissage = COULEURS_CRITICITE.get(v)

        for c in range(1, len(entetes) + 1):
            cellule = ws.cell(row=r, column=c)
            cellule.font = POLICE_CORPS
            cellule.border = BORDURE
            cellule.alignment = Alignment(vertical="top", wrap_text=True)
            if remplissage:
                cellule.fill = PatternFill("solid", fgColor=remplissage)

        # Un indicateur qualifié mais jamais collecté : la distinction est
        # le coeur de l'honnêteté du document, elle doit se voir.
        if col_obs and not ws.cell(row=r, column=col_obs).value:
            cellule = ws.cell(row=r, column=col_obs)
            cellule.value = "aucune collecte"
            cellule.font = Font(name=POLICE, size=10, italic=True, color="8A9199")

    for i, entete in enumerate(entetes, start=1):
        longueurs = [len(str(entete))] + [
            min(len(str(ws.cell(row=r, column=i).value or "")), 120)
            for r in range(2, ws.max_row + 1)
        ]
        plafond = LARGEURS_MAX.get(entete, 38)
        ws.column_dimensions[get_column_letter(i)].width = min(
            max(11, max(longueurs) + 2), plafond
        )

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(entetes))}{ws.max_row}"

    # Mise en page : le classeur doit s'imprimer proprement, il est destiné
    # à circuler en séance de supervision.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"

    return ws


def bloc_de_controle(ws_bilan, ws_indicateurs):
    """Recompte le bilan depuis les lignes d'indicateurs.

    Le décompte qui fait foi est celui de la vue `v_bilan_referentiel`,
    reprise telle quelle dans la feuille Bilan. Ce bloc ne le remplace
    pas : il le recalcule par formule depuis la feuille Indicateurs, de
    sorte qu'un écart entre les deux devienne visible dans le classeur
    lui-même plutôt que de rester silencieux.
    """
    if ws_bilan is None or ws_indicateurs is None:
        return

    entetes = [c.value for c in ws_indicateurs[1]]
    col_statut = index_colonne(entetes, "Statut")
    col_categorie = index_colonne(entetes, "Catégorie")
    if not (col_statut and col_categorie):
        return

    n = ws_indicateurs.max_row
    lettre_statut = get_column_letter(col_statut)
    lettre_cat = get_column_letter(col_categorie)
    plage_statut = f"Indicateurs!${lettre_statut}$2:${lettre_statut}${n}"
    plage_cat = f"Indicateurs!${lettre_cat}$2:${lettre_cat}${n}"

    depart = ws_bilan.max_row + 3
    ws_bilan.cell(row=depart, column=1, value="Contrôle — recalcul par formule").font = Font(
        name=POLICE, bold=True, size=11, color=ENCRE
    )
    ws_bilan.cell(
        row=depart + 1,
        column=1,
        value=(
            "Le décompte ci-dessus provient de la vue v_bilan_referentiel et fait foi. "
            "Les valeurs ci-dessous le recalculent depuis la feuille Indicateurs : "
            "tout écart signale une incohérence à instruire, non une valeur à corriger ici."
        ),
    ).font = Font(name=POLICE, size=9, italic=True, color="5A6570")
    ws_bilan.merge_cells(
        start_row=depart + 1, start_column=1, end_row=depart + 1, end_column=6
    )
    ws_bilan.cell(row=depart + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws_bilan.row_dimensions[depart + 1].height = 30

    controles = [
        ("Total des indicateurs", f'=COUNTA({plage_statut})'),
        ("Certifiés", f'=COUNTIF({plage_statut},"certifie")'),
        (
            "dont hard data",
            f'=COUNTIFS({plage_statut},"certifie",{plage_cat},"hard")',
        ),
        (
            "dont composites",
            f'=COUNTIFS({plage_statut},"certifie",{plage_cat},"composite")',
        ),
        ("À confirmer", f'=COUNTIF({plage_statut},"a_confirmer")'),
    ]
    for decalage, (libelle, formule) in enumerate(controles, start=depart + 3):
        ws_bilan.cell(row=decalage, column=1, value=libelle).font = POLICE_CORPS
        cellule = ws_bilan.cell(row=decalage, column=2, value=formule)
        cellule.font = Font(name=POLICE, size=10, bold=True)
        cellule.border = BORDURE


def feuille_lisez_moi(classeur, horodatage):
    ws = classeur.create_sheet("Lisez-moi", 0)
    lignes = [
        ("Tableau de veille consolidé — annexe 1", ""),
        ("", ""),
        ("Travail de Bachelor", "Exploration de l'IA pour les entreprises industrielles"),
        ("Auteur", "Nilo Castillo — HEG Arc, Bachelor Informatique de gestion"),
        ("Directeur", "Francesco Termine"),
        ("Généré le", horodatage),
        ("", ""),
        ("PROVENANCE — deux natures à ne pas confondre", ""),
        (
            "Le jugement est manuel",
            "Les qualifications consignées ici résultent d'un examen manuel, source par source : accès testé, "
            "contenu vérifié, profondeur d'historique constatée, nomenclature relevée. Statuts de source, "
            "catégories d'indicateur, formulations sectorielles et criticités sont des décisions d'analyse, "
            "prises et assumées par l'auteur. Aucun automatisme ne les produit : la base refuse toute source "
            "dont le qualificateur n'est pas nommé (colonnes « Qualifiée par » et « Le »).",
        ),
        (
            "Le document est généré",
            "Ce classeur, en revanche, est une projection datée du référentiel : il le restitue sans le "
            "ressaisir. Une information saisie deux fois en deux endroits finit toujours par diverger — c'est "
            "ce mécanisme qui avait produit quatre décomptes d'indicateurs contradictoires dans une version "
            "antérieure du rapport.",
        ),
        (
            "Conséquence",
            "Toute divergence entre ce classeur et le corps du rapport signale une erreur du rapport, "
            "non du classeur.",
        ),
        (
            "Régénération",
            "bash exports/generer_csv.sh, puis le conteneur de conversion (voir l'en-tête du script). "
            "À refaire après toute exécution du dispositif.",
        ),
        ("", ""),
        ("STRUCTURE — les sept feuilles", ""),
        (
            "Bilan",
            "Décompte du référentiel, calculé par la vue v_bilan_referentiel. C'est le chiffre qui fait foi ; "
            "il doit être reporté tel quel dans le rapport, jamais l'inverse. Un bloc de contrôle le recalcule "
            "par formule depuis la feuille Indicateurs.",
        ),
        (
            "Indicateurs",
            "La grille complète du chapitre 8, à plat. Filtrable par secteur, statut, catégorie, question de veille. "
            "Porte aussi l'état d'instrumentation de chaque indicateur.",
        ),
        (
            "Instanciation QV",
            "Le niveau 2 du cadre : pour chaque couple secteur × question, la formulation propre au secteur, "
            "le mécanisme causal en jeu et la criticité. C'est la feuille qui distingue le dispositif d'une "
            "simple liste d'indicateurs.",
        ),
        (
            "Questions de veille",
            "Le niveau 1 : les six angles invariants QV0 à QV5, dont cinq sectoriels et un transversal (QV0).",
        ),
        (
            "Couverture QV",
            "Matrice secteur × question. Son intérêt est de rendre l'absence calculable : c'est le cadre "
            "invariant qui permet de constater un vide, un jeu de questions sur mesure définirait sa propre "
            "complétude.",
        ),
        (
            "Lacunes",
            "Extraction des couples sans indicateur certifié, triés par criticité. Ce que le dispositif ne "
            "couvre pas doit se lire aussi vite que ce qu'il couvre.",
        ),
        (
            "Sources",
            "Le tableau de confiance complet. Le § 8.3 du rapport n'en présente que le socle transversal.",
        ),
        ("", ""),
        ("LECTURES À NE PAS MANQUER", ""),
        (
            "Colonne « Observations en base »",
            "Distingue un indicateur QUALIFIÉ d'un indicateur COLLECTÉ. « aucune collecte » signale une "
            "conception non encore instrumentée — limite du travail, énoncée et non masquée.",
        ),
        (
            "Colonne « Catégorie »",
            "hard = donnée officielle structurée reprise telle quelle, traitée par ETL déterministe. "
            "composite = extraite d'un document non structuré, soumise au contrôle de consistance multi-modèles "
            "puis à validation humaine.",
        ),
        (
            "Colonne « Criticité »",
            "Poids de l'angle dans le secteur : dominante, significative ou marginale. Une lacune sur un angle "
            "dominant ne se lit pas comme une lacune sur un angle marginal.",
        ),
        ("", ""),
        (
            "Code couleur",
            "Vert : certifié ou couvert · Ambre : à confirmer ou couverture nominale · Rouge : accès restreint "
            "ou lacune · Violet pâle : criticité, du plus soutenu au plus clair.",
        ),
    ]
    for cle, valeur in lignes:
        ws.append([cle, valeur])

    for r in range(1, ws.max_row + 1):
        cle = ws.cell(row=r, column=1)
        valeur = ws.cell(row=r, column=2)
        est_rubrique = cle.value and not valeur.value and cle.value.isupper()
        cle.font = Font(
            name=POLICE,
            bold=True,
            size=11 if est_rubrique else 10,
            color=ENCRE if est_rubrique else "000000",
        )
        cle.alignment = Alignment(vertical="top", wrap_text=True)
        valeur.font = POLICE_CORPS
        valeur.alignment = Alignment(wrap_text=True, vertical="top")

    titre = ws.cell(row=1, column=1)
    titre.font = Font(name=POLICE, bold=True, size=15, color=ENCRE)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def main():
    horodatage = datetime.datetime.now().strftime("%d.%m.%Y à %H:%M")
    classeur = Workbook()
    classeur.remove(classeur.active)

    feuilles = {}
    for fichier, titre in FEUILLES:
        chemin = CSV / fichier
        if chemin.exists():
            feuilles[titre] = ajouter_feuille(classeur, chemin, titre)
        else:
            print(f"  ! {fichier} absent — feuille « {titre} » ignorée")

    bloc_de_controle(feuilles.get("Bilan"), feuilles.get("Indicateurs"))
    feuille_lisez_moi(classeur, horodatage)

    classeur.save(SORTIE)
    print(f"Classeur écrit : {SORTIE}")
    print(f"Feuilles : {', '.join(classeur.sheetnames)}")


if __name__ == "__main__":
    main()
